"""Export docs/cms/TODO.md to a fillable-looking Word checklist."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "TODO.md"
OUT = HERE / "TODO.docx"
DESKTOP = Path.home() / "Desktop" / "ICBS-Website-TODO.docx"

GREEN = RGBColor(0x1A, 0x3A, 0x32)
INK = RGBColor(0x1C, 0x1C, 0x1C)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)
DONE = RGBColor(0x2E, 0x6B, 0x4A)


def run_font(run, *, size=12, bold=False, italic=False, color=INK):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Calibri")


def add_runs_markdown(paragraph, text, *, size=12, color=INK):
    """Very small subset: **bold**, `code`."""
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            run_font(r, size=size, bold=True, color=GREEN)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            run_font(r, size=size, italic=True, color=INK)
        else:
            cleaned = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", part)
            r = paragraph.add_run(cleaned)
            run_font(r, size=size, color=color)


def para(doc, text, *, size=12, bold=False, color=INK, after=8, before=0, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if bold and "**" not in text:
        r = p.add_run(text)
        run_font(r, size=size, bold=True, color=color, italic=italic)
    else:
        add_runs_markdown(p, text, size=size, color=color)
    return p


def build() -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)

    para(doc, "I Call BS / VettaJimale.Tech", size=12, bold=True, color=GREEN, after=2)
    para(doc, "To-do list (revision)", size=26, bold=True, color=GREEN, after=6)
    para(
        doc,
        "Open this file in Microsoft Word. Check boxes off as items ship. New section: Social launch (SM-01 to SM-08).",
        size=12,
        italic=True,
        color=MUTED,
        after=14,
    )

    for raw in SRC.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.strip() == "---":
            continue
        if line.startswith("# "):
            para(doc, line[2:], size=20, bold=True, color=GREEN, before=4, after=8)
            continue
        if line.startswith("## "):
            para(doc, line[3:], size=16, bold=True, color=GREEN, before=16, after=8)
            continue
        if line.startswith("### "):
            para(doc, line[4:], size=13, bold=True, color=GREEN, before=10, after=6)
            continue
        task = re.match(r"^- \[([ xX])\]\s+(.*)$", line)
        if task:
            done = task.group(1).lower() == "x"
            body = task.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.space_after = Pt(6)
            mark = p.add_run("☑  " if done else "☐  ")
            run_font(mark, size=12, bold=True, color=DONE if done else GREEN)
            add_runs_markdown(p, body, size=12, color=MUTED if done else INK)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.clear()
            add_runs_markdown(p, line[2:], size=12, color=INK)
            p.paragraph_format.space_after = Pt(4)
            continue
        para(doc, line, size=12, color=INK, after=8)

    doc.save(OUT)
    shutil.copy2(OUT, DESKTOP)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
    print(DESKTOP)

"""Export TODO-PLAN.md as a simple numbered Word guide."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "TODO-PLAN.md"
OUT = HERE / "TODO-STEPS.docx"
DESKTOP = Path.home() / "Desktop" / "ICBS-Website-TODO-STEPS.docx"

GREEN = RGBColor(0x1A, 0x3A, 0x32)
INK = RGBColor(0x1C, 0x1C, 0x1C)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)


def font(run, *, size=13, bold=False, italic=False, color=INK):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Calibri")


def clean(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def P(doc, text, *, size=13, bold=False, italic=False, color=INK, after=8, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    r = p.add_run(clean(text))
    font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    r = p.add_run(clean(text))
    font(r, size=13)
    p.paragraph_format.space_after = Pt(4)


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.clear()
    r = p.add_run(clean(re.sub(r"^\d+\.\s*", "", text)))
    font(r, size=13)
    p.paragraph_format.space_after = Pt(4)


def build() -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.9)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)

    P(doc, "I Call BS / VettaJimale.Tech", size=12, bold=True, color=GREEN, after=2)
    P(doc, "To-do plan: numbered steps", size=26, bold=True, color=GREEN, after=8)
    P(
        doc,
        "This is the step-by-step version of the checklist. Open it in Microsoft Word, not in Cursor. Do one task. Follow the steps in order. Check it off in the other Word file (ICBS-Website-TODO.docx) only when the “Done when” test is true.",
        size=14,
        after=10,
    )
    P(doc, "How to use this", size=16, bold=True, color=GREEN, before=6)
    P(doc, "1. Start at Task 1 unless you already finished that item.", size=13)
    P(doc, "2. Read Why (so you know what problem you are solving).", size=13)
    P(doc, "3. Read Done when (so you know when to stop).", size=13)
    P(doc, "4. Do the Steps in order. If a step needs a password or API key you do not have, stop and get access. Do not invent secrets.", size=13)
    P(doc, "5. Then go to the next task.", size=13, after=14)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    task_n = 0
    in_intro = True
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("## "):
            in_intro = False
            title = clean(line[3:])
            if title.lower().startswith("quick"):
                P(doc, title, size=18, bold=True, color=GREEN, before=18, after=8)
                i += 1
                continue
            if title.lower().startswith("after you"):
                P(doc, title, size=18, bold=True, color=GREEN, before=18, after=8)
                i += 1
                continue
            P(doc, title, size=18, bold=True, color=GREEN, before=18, after=8)
            i += 1
            continue
        if line.startswith("### "):
            in_intro = False
            task_n += 1
            P(doc, f"Task {task_n}. {clean(line[4:])}", size=16, bold=True, color=GREEN, before=16, after=8)
            i += 1
            continue
        if in_intro:
            i += 1
            continue
        if not line or line.strip() == "---":
            i += 1
            continue
        if line.startswith("|"):
            i += 1
            continue
        if line.startswith("**Why:**"):
            P(doc, "Why this matters", size=13, bold=True, color=GREEN, before=8, after=2)
            P(doc, line.split(":**", 1)[-1], size=13, after=8)
            i += 1
            continue
        if line.startswith("**Done when:**"):
            P(doc, "Done when", size=13, bold=True, color=GREEN, before=8, after=2)
            P(doc, line.split(":**", 1)[-1], size=13, after=8)
            i += 1
            continue
        if line.startswith("**Status:**"):
            P(doc, "Status", size=13, bold=True, color=GREEN, before=8, after=2)
            P(doc, line.split(":**", 1)[-1], size=13, after=8)
            i += 1
            continue
        if line.startswith("**Steps") or line.strip() == "**Steps**":
            P(doc, "Steps (do these in order)", size=13, bold=True, color=GREEN, before=8, after=4)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            numbered(doc, line)
            i += 1
            continue
        if line.startswith("- "):
            bullet(doc, line[2:])
            i += 1
            continue
        if line.startswith("**") and line.endswith("**"):
            P(doc, line.strip("*"), size=13, bold=True, color=GREEN, before=8, after=4)
            i += 1
            continue
        P(doc, line, size=13, after=8)
        i += 1

    doc.save(OUT)
    shutil.copy2(OUT, DESKTOP)
    return OUT


if __name__ == "__main__":
    print(build())
    print(DESKTOP)

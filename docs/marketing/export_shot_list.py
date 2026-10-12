"""Readable HTML + Desktop Word/PPT for the five-video shot list."""

from __future__ import annotations

import re
import shutil
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.util import Emu, Inches as PInches, Pt as PPt

HERE = Path(__file__).resolve().parent
DESKTOP = Path.home() / "Desktop"
MD = HERE / "SOCIAL-SHOT-LIST.md"
GREEN = RGBColor(0x1A, 0x3A, 0x32)
INK = RGBColor(0x1C, 0x1C, 0x1C)
PPT_GREEN = PptRGB(0x1A, 0x3A, 0x32)
PPT_CREAM = PptRGB(0xF7, 0xF2, 0xE8)
PPT_WHITE = PptRGB(0xFF, 0xFF, 0xFF)
PPT_INK = PptRGB(0x1C, 0x1C, 0x1C)


def md_to_html(src: str) -> str:
    lines = src.splitlines()
    out: list[str] = [
        "<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'>",
        "<meta name='viewport' content='width=device-width, initial-scale=1'>",
        "<title>Five-video shot list</title><style>",
        "body{margin:0;background:#f7f2e8;color:#1c1c1c;font:20px/1.45 Calibri,Segoe UI,sans-serif;}",
        "main{max-width:52rem;margin:0 auto;padding:2rem 1.25rem 4rem;}",
        "h1{font-size:2rem;color:#1a3a32;}h2{font-size:1.5rem;color:#1a3a32;margin-top:2.4rem;}",
        "h3{font-size:1.2rem;color:#1a3a32;}hr{border:0;border-top:2px solid #1a3a32;margin:2rem 0;}",
        "table{border-collapse:collapse;width:100%;font-size:16px;margin:1rem 0 1.4rem;background:#fff;}",
        "th,td{border:1px solid #1a3a32;padding:.55rem .65rem;text-align:left;vertical-align:top;}",
        "th{background:#1a3a32;color:#f7f2e8;}strong{color:#1a3a32;}",
        "ul,ol{padding-left:1.2rem;}li{margin:.25rem 0;}",
        "</style></head><body><main>",
    ]

    def inline(t: str) -> str:
        t = escape(t)
        t = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)
        t = re.sub(r"`(.+?)`", r"<code>\1</code>", t)
        return t

    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("| ") and i + 1 < len(lines) and re.match(r"^\|[\s\-|]+\|$", lines[i + 1] or ""):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s\-|]+\|$", lines[i]):
                    cells = [c.strip() for c in lines[i].strip("|").split("|")]
                    rows.append(cells)
                i += 1
            out.append("<table>")
            for r, row in enumerate(rows):
                tag = "th" if r == 0 else "td"
                out.append("<tr>" + "".join(f"<{tag}>{inline(c)}</{tag}>" for c in row) + "</tr>")
            out.append("</table>")
            continue
        if line.startswith("# "):
            out.append(f"<h1>{inline(line[2:])}</h1>")
        elif line.startswith("## "):
            out.append(f"<h2>{inline(line[3:])}</h2>")
        elif line.startswith("### "):
            out.append(f"<h3>{inline(line[4:])}</h3>")
        elif line.strip() == "---":
            out.append("<hr>")
        elif line.startswith("- "):
            out.append("<ul>")
            while i < len(lines) and lines[i].startswith("- "):
                out.append(f"<li>{inline(lines[i][2:])}</li>")
                i += 1
            out.append("</ul>")
            continue
        elif re.match(r"^\d+\. ", line):
            out.append("<ol>")
            while i < len(lines) and re.match(r"^\d+\. ", lines[i]):
                out.append(f"<li>{inline(re.sub(r'^\d+\. ', '', lines[i]))}</li>")
                i += 1
            out.append("</ol>")
            continue
        elif line.strip():
            out.append(f"<p>{inline(line)}</p>")
        i += 1
    out.append("</main></body></html>")
    return "".join(out)


def set_run(run, *, size=14, bold=False, color=INK):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Calibri")


def build_docx(text: str, dest: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)
    for raw in text.splitlines():
        line = raw.rstrip()
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        if line.startswith("# "):
            set_run(p.add_run(line[2:]), size=22, bold=True, color=GREEN)
        elif line.startswith("## "):
            set_run(p.add_run(line[3:]), size=16, bold=True, color=GREEN)
        elif line.startswith("### "):
            set_run(p.add_run(line[4:]), size=14, bold=True, color=GREEN)
        else:
            set_run(p.add_run(line), size=13, color=INK)
    doc.save(dest)


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _write(slide, l, t, w, h, lines, *, size=20, bold=False, color=PPT_INK):
    box = slide.shapes.add_textbox(PInches(l), PInches(t), PInches(w), PInches(h))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = line
        p.font.name = "Calibri"
        p.font.size = PPt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = PPt(8)


def build_pptx(dest: Path) -> None:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)

    s = _blank(prs)
    _bg(s, PPT_GREEN)
    _write(s, 0.7, 2.0, 12, 1, ["VettaJimale.Tech"], size=18, bold=True, color=PPT_CREAM)
    _write(s, 0.7, 2.6, 12, 2, ["Five-video shot list"], size=36, bold=True, color=PPT_WHITE)
    _write(s, 0.7, 5.0, 12, 1.4, ["Scripts 2, 5, 10, 11, 15. Film in batch order. Post over three weeks."], size=20, color=PPT_CREAM)

    s = _blank(prs)
    _bg(s, PPT_CREAM)
    _write(s, 0.6, 0.4, 12, 0.6, ["Batch order (one afternoon)"], size=26, bold=True, color=PPT_GREEN)
    _write(
        s,
        0.6,
        1.2,
        12,
        5.6,
        [
            "A  Silent B-roll — 20 min",
            "B  Script 2 spelling — 15–25s",
            "C  Script 10 literacy check — 40–55s",
            "D  Script 11 templates — 70–90s",
            "E  Script 5 three tools — 45–70s",
            "F  Script 15 pin — 90–120s",
        ],
        size=22,
        color=PPT_INK,
    )

    s = _blank(prs)
    _bg(s, PPT_CREAM)
    _write(s, 0.6, 0.4, 12, 0.6, ["Post over three weeks"], size=26, bold=True, color=PPT_GREEN)
    _write(
        s,
        0.6,
        1.2,
        12,
        5.6,
        [
            "Week 1 Mon   2  Spelling",
            "Week 1 Thu  10  Literacy check",
            "Week 2 Mon  11  Templates",
            "Week 2 Thu   5  Three live tools",
            "Week 3 Mon  15  Pin this one",
            "Bumper: last 3 seconds of #2 whenever you say the name.",
        ],
        size=22,
        color=PPT_INK,
    )

    videos = [
        ("2  Spelling", "15–25s", "Last frame = URL card for 3 seconds. Thumbnail is the spelled URL, not your face."),
        ("10  Literacy check", "40–55s", "Show one question + explanation. Never the whole quiz. No score-for-clout ending."),
        ("11  Templates", "70–90s", "Fill real lines. Show Markdown and PDF. Then prompt."),
        ("5  Three live tools", "45–70s", "Reuse screen from 10 and 11. Fresh pass on Tools & Updates + limits."),
        ("15  The pin", "90–120s", "Say the empty shelf. Hold VettaJimale.Tech for 8 seconds. Pin it. Do not replace with a trailer."),
    ]
    for title, runtime, note in videos:
        s = _blank(prs)
        _bg(s, PPT_CREAM)
        _write(s, 0.6, 0.4, 12, 0.5, [f"Video {title}"], size=26, bold=True, color=PPT_GREEN)
        _write(s, 0.6, 1.2, 12, 0.5, [runtime], size=20, bold=True, color=PPT_GREEN)
        _write(s, 0.6, 2.0, 12, 4, [note, "Full run of show, comments, and captions: the Word/HTML shot list."], size=22, color=PPT_INK)

    prs.save(dest)


def main() -> None:
    text = MD.read_text(encoding="utf-8")
    html_path = HERE / "VettaJimale-Tech-Shot-List.html"
    html_path.write_text(md_to_html(text), encoding="utf-8")
    DESKTOP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(MD, DESKTOP / "VettaJimale-Tech-Shot-List.md")
    shutil.copy2(html_path, DESKTOP / html_path.name)
    build_docx(text, DESKTOP / "VettaJimale-Tech-Shot-List.docx")
    build_pptx(DESKTOP / "VettaJimale-Tech-Shot-List.pptx")
    print("html", html_path)
    print("desktop copies written")


if __name__ == "__main__":
    main()
